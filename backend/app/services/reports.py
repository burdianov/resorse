"""Building reports: a PDF renderer, a DOCX renderer and a merge (F051/F052, BP-7.9b).

Three operations, and no route. F053 owns the endpoint, the ``reports.generate``
check and the preview UI. What lives here is the part no caller should
re-implement: a ReportLab document that carries the furniture a report needs — a
title block, a metadata block, tables whose headers repeat when they spill — the
same :class:`ReportDocument` rendered as a Word file through docxtpl, and a pypdf
merge that puts several documents in one file.

Three positions worth naming, because each is a place the obvious shortcut is
wrong:

- **The page count is real, not guessed.** "Page 1 of 3" is only a true
  sentence if the total is known before the first page is drawn, so the canvas
  keeps its page states and stamps the footer in ``save()``. A document that
  grows a page because a table spilled therefore still tells the truth on
  page one.
- **Only the fonts ReportLab already carries.** Helvetica needs no font file to
  ship, keeps the bytes small, and keeps the text *extractable* — which is what
  lets the acceptance ("the generated PDF parses") be checked by reading the
  words back out. The cost is stated rather than hidden: these fonts cover
  Latin-1, so a script that needs glyph shaping (Arabic, CJK) renders as blanks
  until an embedded TTF is chosen for it, and that is a decision for the first
  task that needs one (C39).
- **A value is a value.** ``<`` or ``&`` in a cell is text a user typed, not
  markup: every string that reaches a ``Paragraph`` is escaped first, and the
  DOCX renderer switches on docxtpl's own ``autoescape`` for the same reason
  (escaping by hand *there* would print ``&amp;amp;``). BP-7.9d's CSV/Excel half
  calls the same class of problem formula injection; the PDF's version is markup
  injection, and it is closed at this boundary rather than trusted to every call
  site.

Nothing here reads settings or the database. The caller supplies the words and
the data, so the engine holds no opinion about branding (C29 is still open) and
can be exercised without a session.

The functions are **synchronous**: ReportLab is CPU-bound and does no I/O, so
there is nothing to await. An endpoint rendering a large report should hand it
to a worker thread rather than block the event loop — F053's decision, made
where it owns the route.
"""

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from enum import StrEnum
from functools import lru_cache, partial
from typing import Any, Final
from xml.sax.saxutils import escape

from docx import Document as DocxDocument
from docxtpl import DocxTemplate
from pypdf import PageObject, PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class ReportError(Exception):
    """Something this engine refuses to render or read."""


class UnreadablePdf(ReportError):
    """The bytes are not a PDF, or are damaged past reading.

    Raised instead of pypdf's own exception: whatever the library happened to
    raise — and it has several, for several kinds of damage — the caller's
    situation is one thing, and a caller that has to name the library's error
    class to handle it has been handed an implementation detail.
    """


class EmptyMerge(ReportError):
    """A merge was asked for with nothing to merge.

    A zero-page PDF is not an empty bundle; it is a document that lies about
    being one, so the refusal happens here rather than three pages later in a
    viewer that renders a blank sheet.
    """


class PageLayout(StrEnum):
    """The four geometries the stack's reports use (CLAUDE_MASTER §Engineering).

    A caller picks a name, never a tuple of points: A3 landscape is the default
    for wide manpower matrices and A4 for documents, and both are here so the
    choice is explicit at the call site rather than buried in a geometry.
    """

    A4_PORTRAIT = "a4_portrait"
    A4_LANDSCAPE = "a4_landscape"
    A3_PORTRAIT = "a3_portrait"
    A3_LANDSCAPE = "a3_landscape"

    @property
    def page_size(self) -> tuple[float, float]:
        """The geometry as ``(width, height)`` in points."""
        return _PAGE_SIZES[self]


_PAGE_SIZES: Final[dict[PageLayout, tuple[float, float]]] = {
    PageLayout.A4_PORTRAIT: A4,
    PageLayout.A4_LANDSCAPE: landscape(A4),
    PageLayout.A3_PORTRAIT: A3,
    PageLayout.A3_LANDSCAPE: landscape(A3),
}


@dataclass(frozen=True)
class ReportSection:
    """A heading and the paragraphs under it."""

    heading: str
    paragraphs: Sequence[str] = ()


@dataclass(frozen=True)
class ReportTable:
    """A header row and rectangular data, rendered as strings.

    Rows are *not* trusted to be rectangular: a short row is a caller's bug,
    and ReportLab would answer it with a layout that quietly drops cells, so
    :func:`render_report` refuses it by name instead (see :func:`_checked`).
    """

    columns: Sequence[str]
    rows: Sequence[Sequence[Any]] = ()
    heading: str | None = None


@dataclass(frozen=True)
class ReportDocument:
    """Everything a report says, in the order it says it.

    ``metadata`` is the block of label/value facts an executive report opens
    with (period, revision, scope — the words belong to the caller); the
    document's *own* metadata (title, author, subject, generated-at) travels
    into the PDF's information dictionary and its footer, so a file that leaves
    this service still says where it came from.
    """

    title: str
    subtitle: str | None = None
    metadata: Sequence[tuple[str, str]] = ()
    sections: Sequence[ReportSection] = ()
    tables: Sequence[ReportTable] = ()
    layout: PageLayout = PageLayout.A4_PORTRAIT
    author: str = "resors"
    subject: str | None = None
    keywords: Sequence[str] = ()
    generated_at: datetime | None = None


# --------------------------------------------------------------------------
# Layout constants and styles
# --------------------------------------------------------------------------

MARGIN: Final = 18 * mm
TOP_MARGIN: Final = 16 * mm
# The footer rule sits at 15.5mm; the frame has to stop above it.
BOTTOM_MARGIN: Final = 20 * mm
FOOTER_BASELINE: Final = 12 * mm
FOOTER_FONT: Final = "Helvetica"
FOOTER_SIZE: Final = 7.5
CREATOR: Final = "resors"

_RULE = colors.HexColor("#b8b8b8")
_HEADER_FILL = colors.HexColor("#ededed")
_QUIET = colors.HexColor("#555555")

_BASE = getSampleStyleSheet()

_TITLE = ParagraphStyle(
    "ReportTitle",
    parent=_BASE["Title"],
    fontName="Helvetica-Bold",
    fontSize=17,
    leading=21,
    alignment=0,
    spaceAfter=0,
)
_SUBTITLE = ParagraphStyle(
    "ReportSubtitle",
    parent=_BASE["Normal"],
    fontName="Helvetica",
    fontSize=10.5,
    leading=14,
    textColor=_QUIET,
    spaceBefore=2,
)
_HEADING = ParagraphStyle(
    "ReportHeading",
    parent=_BASE["Normal"],
    fontName="Helvetica-Bold",
    fontSize=11.5,
    leading=15,
    spaceBefore=10,
    spaceAfter=3,
)
_BODY = ParagraphStyle(
    "ReportBody",
    parent=_BASE["Normal"],
    fontName="Helvetica",
    fontSize=9.5,
    leading=12.5,
)
_LABEL = ParagraphStyle(
    "ReportLabel",
    parent=_BASE["Normal"],
    fontName="Helvetica-Bold",
    fontSize=8.5,
    leading=11.5,
    textColor=_QUIET,
)
_VALUE = ParagraphStyle(
    "ReportValue",
    parent=_BASE["Normal"],
    fontName="Helvetica",
    fontSize=9,
    leading=11.5,
)
_CELL = ParagraphStyle(
    "ReportCell",
    parent=_BASE["Normal"],
    fontName="Helvetica",
    fontSize=8.5,
    leading=10.5,
)
_COLUMN = ParagraphStyle(
    "ReportColumn",
    parent=_BASE["Normal"],
    fontName="Helvetica-Bold",
    fontSize=8.5,
    leading=10.5,
)


def escaped(value: Any) -> str:
    """``value`` as text a ``Paragraph`` renders literally.

    Public because the escaping rule is the engine's contract with its callers,
    not an internal detail: any string handed to a report goes through here, so
    a workforce named ``R&D <night>`` prints as typed.
    """
    return escape(str(value))


def _utc_stamp(moment: datetime) -> str:
    """An instant as both renderers write it: UTC, to the minute, labelled.

    Display conversion is the settings/UI layer's business (F048); a file is a
    record, and a record states the zone it was written in.
    """
    return f"{moment.astimezone(UTC):%Y-%m-%d %H:%M} UTC"


# --------------------------------------------------------------------------
# Rendering
# --------------------------------------------------------------------------


def render_report(document: ReportDocument) -> bytes:
    """Render ``document`` to the bytes of a complete PDF.

    The returned file carries the document's own metadata in its information
    dictionary and, on every page, a footer naming the report, the moment it
    was generated (UTC — display conversion is the settings/UI layer's job,
    C29/F048) and the page number out of the real total.
    """
    generated_at = document.generated_at or datetime.now(UTC)
    buffer = io.BytesIO()
    document_template = SimpleDocTemplate(
        buffer,
        pagesize=document.layout.page_size,
        leftMargin=MARGIN,
        rightMargin=MARGIN,
        topMargin=TOP_MARGIN,
        bottomMargin=BOTTOM_MARGIN,
        title=document.title,
        author=document.author,
        subject=document.subject or document.title,
        creator=CREATOR,
        keywords=list(document.keywords) or None,
    )
    document_template.build(
        _story(document),
        canvasmaker=partial(
            _ReportCanvas,
            footer_left=document.title,
            footer_right=_utc_stamp(generated_at),
        ),
    )
    return buffer.getvalue()


def _story(document: ReportDocument) -> list[Flowable]:
    """The document as flowables, in the order a reader meets them."""
    story: list[Flowable] = [Paragraph(escaped(document.title), _TITLE)]
    if document.subtitle:
        story.append(Paragraph(escaped(document.subtitle), _SUBTITLE))
    if document.metadata:
        story.append(Spacer(1, 5 * mm))
        story.append(_fact_table(document.metadata))
    for section in document.sections:
        story.append(Paragraph(escaped(section.heading), _HEADING))
        for paragraph in section.paragraphs:
            story.append(Paragraph(escaped(paragraph), _BODY))
    for report_table in document.tables:
        if report_table.heading:
            story.append(Paragraph(escaped(report_table.heading), _HEADING))
        story.append(_data_table(report_table))
        story.append(Spacer(1, 3 * mm))
    return story


def _fact_table(metadata: Sequence[tuple[str, str]]) -> Table:
    """The label/value block, drawn as a borderless two-column table."""
    rows = [
        [Paragraph(escaped(label), _LABEL), Paragraph(escaped(value), _VALUE)]
        for label, value in metadata
    ]
    # A `None` column width takes what is left of the frame, so the values
    # align without this function knowing the page width.
    table = Table(rows, colWidths=[26 * mm, None], hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    return table


def _data_table(report_table: ReportTable) -> Table:
    """A report table: header row, data rows, and a header that repeats."""
    columns, rows = _checked(report_table)
    data: list[list[Paragraph]] = [[Paragraph(escaped(column), _COLUMN) for column in columns]]
    data.extend([Paragraph(escaped(cell), _CELL) for cell in row] for row in rows)
    table = Table(data, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, _RULE),
                ("BACKGROUND", (0, 0), (-1, 0), _HEADER_FILL),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def _checked(report_table: ReportTable) -> tuple[Sequence[str], Sequence[Sequence[Any]]]:
    """The table's columns and rows, or a :class:`ReportError` naming the fault.

    A table is either rectangular or it is a mistake: ReportLab would render a
    short row by dropping cells silently, which is the one outcome a report
    must never have (a number that is missing looks like a number that is
    zero).
    """
    columns = report_table.columns
    if not columns:
        raise ReportError("A report table needs at least one column.")
    for index, row in enumerate(report_table.rows, start=1):
        if len(row) != len(columns):
            raise ReportError(
                f"Row {index} has {len(row)} cells and the table has "
                f"{len(columns)} columns. A report table is rectangular."
            )
    return columns, report_table.rows


def _fit(text: str, *, width: float, font: str, size: float) -> str:
    """``text`` shortened to ``width`` points, marked with a plain ellipsis.

    Three ASCII dots rather than U+2026: the footer is furniture, and furniture
    should not depend on a glyph the built-in fonts may or may not carry.
    """
    if stringWidth(text, font, size) <= width:
        return text
    trimmed = text
    while trimmed and stringWidth(f"{trimmed}...", font, size) > width:
        trimmed = trimmed[:-1]
    return f"{trimmed}..." if trimmed else ""


# ReportLab publishes no stubs (see the override in pyproject.toml), so its
# classes reach a type checker as `Any`, and strict mode refuses to subclass
# `Any` unless the one line that does it says so. Remove this if ReportLab ever
# ships a `py.typed` — mypy will report the ignore as unused the moment it can
# type the base class.
class _ReportCanvas(canvas.Canvas):  # type: ignore[misc]
    """A canvas that numbers each page *and* counts them.

    The page total cannot be known while the first page is being drawn, so the
    pages are drawn to this canvas and held: :meth:`showPage` keeps each page's
    state, and :meth:`save` replays them with the footer, by which time the
    count is a fact. This is ReportLab's own recipe, and it is the reason the
    footer on page one can say "1 of 7".
    """

    def __init__(self, *args: Any, footer_left: str, footer_right: str, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._footer_left = footer_left
        self._footer_right = footer_right
        self._page_states: list[dict[str, Any]] = []

    def showPage(self) -> None:
        self._page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        total = len(self._page_states)
        for state in self._page_states:
            self.__dict__.update(state)
            self._draw_footer(total)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def _draw_footer(self, total: int) -> None:
        width, _ = self._pagesize
        rule = FOOTER_BASELINE + 3.5 * mm
        self.saveState()
        self.setStrokeColor(_RULE)
        self.setLineWidth(0.4)
        self.line(MARGIN, rule, width - MARGIN, rule)
        self.setFont(FOOTER_FONT, FOOTER_SIZE)
        self.setFillColor(colors.black)
        # Half the frame, less the space the centred page number needs.
        self.drawString(
            MARGIN,
            FOOTER_BASELINE,
            _fit(
                self._footer_left,
                width=(width - 2 * MARGIN) * 0.45,
                font=FOOTER_FONT,
                size=FOOTER_SIZE,
            ),
        )
        self.drawCentredString(width / 2, FOOTER_BASELINE, f"Page {self._pageNumber} of {total}")
        self.drawRightString(width - MARGIN, FOOTER_BASELINE, self._footer_right)
        self.restoreState()


# --------------------------------------------------------------------------
# Word documents
# --------------------------------------------------------------------------

# A Word file's type on the wire, named once so the renderer, the converter and
# their callers cannot disagree about it.
DOCX_CONTENT_TYPE: Final = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def render_docx(document: ReportDocument, *, template: bytes | None = None) -> bytes:
    """Render ``document`` to the bytes of a Word file (BP-7.9b's docxtpl half).

    The same document in a different medium: :func:`render_report` is the
    executive PDF — tables, page furniture, A3 landscape — while this is the road
    a report takes to Word and, through Gotenberg (F052), to a PDF as well. It is
    the *narrative* half of the pair: the built-in template draws the title, the
    fact block and the sections, and deliberately not the tables, because a wide
    manpower matrix is the PDF engine's job. A caller that wants tables in Word
    passes ``template``; the context carries them either way
    (:func:`_docx_context`), which is the seam a house template needs.

    ``template`` is the bytes of a ``.docx`` whose text may hold docxtpl tags;
    the default is described in :func:`_default_docx_template`.
    """
    source = template if template is not None else _default_docx_template()
    docx = DocxTemplate(io.BytesIO(source))
    docx.render(_docx_context(document), autoescape=True)
    output = io.BytesIO()
    docx.save(output)
    return output.getvalue()


@lru_cache
def _default_docx_template() -> bytes:
    """The built-in template: a Word file whose text carries the tags.

    Built in code rather than committed as a binary asset. A ``.docx`` nobody can
    read in a diff is exactly what a template should not be; while the default
    lives here, everything it says is visible in review, and a deployment that
    wants its own letterhead passes one in.

    Only two kinds of tag appear, and each is used the way docxtpl requires:
    inline ``{{ … }}`` inside a paragraph, and paragraph-level ``{%p … %}`` alone
    in its own paragraph — the tag stands for the whole paragraph in the XML, so
    anything else sharing that paragraph would be discarded with it.
    """
    template = DocxDocument()
    template.add_heading("{{ title }}", level=0)
    template.add_paragraph("{%p if subtitle %}")
    template.add_paragraph("{{ subtitle }}")
    template.add_paragraph("{%p endif %}")
    template.add_paragraph("{%p for fact in facts %}")
    template.add_paragraph("{{ fact.label }}: {{ fact.value }}")
    template.add_paragraph("{%p endfor %}")
    template.add_paragraph("{%p for section in sections %}")
    template.add_heading("{{ section.heading }}", level=1)
    template.add_paragraph("{%p for paragraph in section.paragraphs %}")
    template.add_paragraph("{{ paragraph }}")
    template.add_paragraph("{%p endfor %}")
    template.add_paragraph("{%p endfor %}")
    template.add_paragraph("{{ author }} · {{ generated_at }}")
    buffer = io.BytesIO()
    template.save(buffer)
    return buffer.getvalue()


def _docx_context(document: ReportDocument) -> dict[str, Any]:
    """The document as the template's context — every word it has, escaped once.

    The values are **raw**: ``autoescape=True`` at the render call is docxtpl's
    own switch, and escaping here as well would double it into ``&amp;amp;``. The
    PDF engine escapes by hand only because ReportLab offers no equivalent.

    Everything a :class:`ReportDocument` can say is present, including the tables
    the built-in template does not draw, so a caller's template is not bounded by
    the one that ships.
    """
    generated_at = document.generated_at or datetime.now(UTC)
    return {
        "title": document.title,
        "subtitle": document.subtitle,
        "author": document.author,
        "generated_at": _utc_stamp(generated_at),
        "facts": [{"label": label, "value": value} for label, value in document.metadata],
        "sections": [
            {"heading": section.heading, "paragraphs": list(section.paragraphs)}
            for section in document.sections
        ],
        "tables": [_docx_table(table) for table in document.tables],
    }


def _docx_table(report_table: ReportTable) -> dict[str, Any]:
    """A report table as context a template can loop over.

    Checked by :func:`_checked` for the same reason the PDF path is: a ragged row
    is a caller's bug, and a document that quietly drops cells is worse than one
    that refuses to be built — a missing number reads as a zero.
    """
    columns, rows = _checked(report_table)
    return {
        "heading": report_table.heading,
        "columns": list(columns),
        "rows": [[str(cell) for cell in row] for row in rows],
    }


# --------------------------------------------------------------------------
# Merging
# --------------------------------------------------------------------------


def merge_pdfs(parts: Sequence[bytes], *, title: str | None = None) -> bytes:
    """Concatenate PDFs — in the order given — into one document.

    Order is the contract: a bundle is a list, not a set, and the caller's
    sequence is the order the reader will meet the pages in. Every part is
    read before anything is written, so a damaged part refuses the merge
    instead of producing a document that is missing its middle.
    """
    if not parts:
        raise EmptyMerge("A merge needs at least one document.")
    pages: list[PageObject] = []
    for index, part in enumerate(parts, start=1):
        pages.extend(_pages(part, where=f"part {index}"))
    writer = PdfWriter()
    for page in pages:
        writer.add_page(page)
    if title:
        writer.add_metadata({"/Title": title})
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def count_pages(pdf: bytes) -> int:
    """How many pages ``pdf`` has, or :class:`UnreadablePdf`."""
    return len(_pages(pdf, where="the document"))


def _pages(pdf: bytes, *, where: str) -> list[PageObject]:
    """``pdf``'s pages, with every way of failing to read them made one failure.

    ``list(...)`` rather than ``len(reader.pages)``: pypdf resolves pages
    lazily, so walking the tree here moves a damaged page object's failure
    inside this boundary instead of leaving it to surface at a caller's
    iteration. The ``except`` is deliberately broad for the same reason — pypdf
    raises several shapes for several kinds of damage, and from outside they
    are one situation.
    """
    try:
        return list(PdfReader(io.BytesIO(pdf)).pages)
    except Exception as broken:
        raise UnreadablePdf(f"{where} is not a PDF this engine can read: {broken}") from broken
